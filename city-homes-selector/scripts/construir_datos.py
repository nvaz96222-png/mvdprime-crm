"""
construir_datos.py — genera src/datos/unidades.json a partir de los documentos originales.

Fuentes (en datos-fuente/):
  · LISTA-OFICIAL-CITY-HOMES.pdf   — lista oficial, Drive > City Homes (mod. 2026-06-24)
  · LISTAS-PIRES-BENLIAN.xlsx      — hoja "CITY HOMES" (mod. 2026-10-23, la mas reciente)

Regla de precedencia, explicita y no arbitraria: cuando las dos listas difieren,
vale la MAS RECIENTE (la planilla Pires Benlian, columna L). La otra lectura se
conserva en `historico` y la diferencia queda registrada en `discrepancias`.

Uso:  python -I scripts/construir_datos.py
"""

import csv
import json
import re
import unicodedata
from pathlib import Path

import openpyxl
import pymupdf

RAIZ = Path(__file__).resolve().parent.parent
FUENTE = RAIZ / "datos-fuente"
DESTINO = RAIZ / "src" / "datos"

PDF = FUENTE / "LISTA-OFICIAL-CITY-HOMES.pdf"
XLSX = FUENTE / "LISTAS-PIRES-BENLIAN.xlsx"

DOC_PDF = "LISTA OFICIAL CITY HOMES.pdf (Drive > City Homes, 2026-06-24)"
DOC_XLSX = 'LISTAS PIRES BENLIAN.xlsx, hoja "CITY HOMES", columna PRECIO vigente (2026-10-23)'


def num(texto):
    """'37,44' -> 37.44 ; '' -> None"""
    if texto in (None, ""):
        return None
    if isinstance(texto, (int, float)):
        return round(float(texto), 2)
    t = str(texto).strip().replace(".", "").replace(",", ".")
    try:
        return round(float(t), 2)
    except ValueError:
        return None


def sin_tildes(t):
    return "".join(c for c in unicodedata.normalize("NFD", t) if unicodedata.category(c) != "Mn")


# ─────────────────────────────────────────────────────────── lista oficial (PDF)

def leer_pdf():
    doc = pymupdf.open(PDF)
    lineas = []
    for pagina in doc:
        lineas += [l.strip() for l in pagina.get_text().splitlines() if l.strip()]

    unidades, garajes = {}, {}
    i, en_garajes = 0, False
    while i < len(lineas):
        linea = lineas[i]
        if linea.upper().startswith("GARAJES"):
            en_garajes = True
            i += 1
            continue
        if en_garajes:
            if re.fullmatch(r"\d{1,2}", linea) and i + 1 < len(lineas):
                garajes[linea] = lineas[i + 1]
                i += 2
                continue
            i += 1
            continue
        m = re.fullmatch(r"AP\.\s*(\d{3,4})", linea)
        if m and i + 7 < len(lineas):
            # 7 valores por unidad: dorm · habitable · terrazas · comunes · % · total · precio
            campos = lineas[i + 1 : i + 8]
            unidades[m.group(1)] = {
                "dormitorios": campos[0],
                "habitable": num(campos[1]),
                "terrazas": num(campos[2]),
                "areasComunes": num(campos[3]),
                "porcentaje": num(campos[4]),
                "total": num(campos[5]),
                "precioBruto": campos[6],
            }
            i += 8
            continue
        i += 1
    return unidades, garajes


# ────────────────────────────────────────────────────── planilla vigente (XLSX)

def leer_xlsx():
    hoja = openpyxl.load_workbook(XLSX, data_only=True)["CITY HOMES"]
    unidades, garajes = {}, {}
    nivel_actual, en_garajes = None, False

    for fila in hoja.iter_rows(min_row=1, max_row=hoja.max_row, max_col=12, values_only=True):
        # `celdas` sólo para texto; los números se leen crudos de `fila` para no
        # pasarlos por el parser de miles/decimales castellano.
        celdas = ["" if c is None else str(c).strip() for c in fila]
        a, b = celdas[0], celdas[1]

        if sin_tildes(a).upper().startswith("GARAJES"):
            en_garajes = True
        if a and "PISO" in sin_tildes(a).upper():
            nivel_actual = a

        if en_garajes:
            if re.fullmatch(r"\d{1,2}", b):
                garajes[b] = (
                    fila[11] if isinstance(fila[11], (int, float)) else (celdas[11] or celdas[10])
                )
            continue

        m = re.fullmatch(r"AP\.\s*(\d{3,4})", b)
        if not m:
            continue
        precio_vigente = fila[11] if isinstance(fila[11], (int, float)) else celdas[11]
        precio_previo = fila[10] if isinstance(fila[10], (int, float)) else celdas[10]
        if not precio_vigente and not precio_previo:
            continue  # fila duplicada de encabezado/subtotal sin precio
        unidades[m.group(1)] = {
            "nivel": nivel_actual,
            "dormitorios": celdas[2],
            "banos": fila[3],
            "habitable": num(fila[4]),
            "terrazas": num(fila[5]),
            "subtotal": num(fila[6]),
            "areasComunes": num(fila[7]),
            "porcentaje": num(fila[8]),
            "total": num(fila[9]),
            "precioPrevio": precio_previo,
            "precioVigente": precio_vigente,
        }
    return unidades, garajes


# ──────────────────────────────────────────────────────────────────── normaliza

def interpretar_precio(bruto):
    """Devuelve (estado, monto|None) a partir del contenido de la celda de precio."""
    if bruto is None or bruto == "":
        return "consultar", None
    if isinstance(bruto, (int, float)):
        return "disponible", int(round(bruto))
    t = str(bruto).strip().upper()
    if "VENDIDO" in t:
        return "vendido", None
    if "RESERVADO" in t:
        return "reservado", None
    monto = num(t.replace("USD", "").strip())
    if monto is None:
        return "consultar", None
    return "disponible", int(round(monto))


TIPOLOGIAS = {
    "MA": ("Monoambiente", 0),
    "1": ("1 dormitorio", 1),
    "2": ("2 dormitorios", 2),
}

PLANO_POR_PISO = {
    1: "piso-01",
    **{p: "pisos-02-09" for p in range(2, 10)},
    10: "piso-10",
}

NIVEL_ETIQUETA = {
    1: "1º PISO", 2: "2º PISO", 3: "3er. PISO", 4: "4º PISO", 5: "5º PISO",
    6: "6º PISO", 7: "7º PISO", 8: "8º PISO", 9: "9º PISO", 10: "10º PISO",
}


def construir():
    pdf_unidades, pdf_garajes = leer_pdf()
    xls_unidades, xls_garajes = leer_xlsx()

    numeros = sorted(set(pdf_unidades) | set(xls_unidades), key=lambda n: (len(n), int(n)))
    unidades, discrepancias = [], []

    for numero in numeros:
        p, x = pdf_unidades.get(numero), xls_unidades.get(numero)
        base = x or p
        piso = int(numero[:-2])
        plano = PLANO_POR_PISO.get(piso)
        zona = numero if piso in (1, 10) else numero[-2:]

        clave_dorm = (base.get("dormitorios") or "").strip().upper()
        tipologia, dormitorios = TIPOLOGIAS.get(clave_dorm, ("Sin especificar", None))

        if x:
            estado, monto = interpretar_precio(x["precioVigente"] or x["precioPrevio"])
            fuente_estado = DOC_XLSX
        else:
            estado, monto = interpretar_precio(p["precioBruto"])
            fuente_estado = DOC_PDF

        estado_pdf = interpretar_precio(p["precioBruto"]) if p else (None, None)

        notas = []
        if p and x:
            if estado_pdf[0] != estado:
                discrepancias.append({
                    "tipo": "estado", "unidad": numero,
                    "listaOficialPdf": estado_pdf[0], "planillaVigente": estado,
                })
                notas.append(
                    f"Estado a confirmar: la lista oficial en PDF (24/06/2026) la da como «{estado_pdf[0]}» "
                    f"y la planilla del 23/10/2026 como «{estado}». Se muestra la más reciente."
                )
            elif estado == "disponible" and estado_pdf[1] and monto and estado_pdf[1] != monto:
                discrepancias.append({
                    "tipo": "precio", "unidad": numero,
                    "listaOficialPdf": estado_pdf[1], "planillaVigente": monto,
                })

        superficies = {
            "habitable": base.get("habitable"),
            "terrazasUsoExclusivo": base.get("terrazas"),
            "subTotal": (x or {}).get("subtotal"),
            "areasComunes": base.get("areasComunes"),
            "total": base.get("total"),
        }

        unidades.append({
            "id": f"ap-{numero}",
            "unidad": numero,
            "etiqueta": f"AP. {numero}",
            "piso": piso,
            "nivel": (x or {}).get("nivel") or NIVEL_ETIQUETA.get(piso, f"{piso}º PISO"),
            "tipologia": tipologia,
            "dormitorios": dormitorios,
            "banos": int(num(x["banos"])) if x and num(x.get("banos")) is not None else None,
            "superficies": superficies,
            "precio": monto,
            "moneda": "USD",
            "estado": estado,
            "plano": plano,
            "zona": zona,
            "observaciones": notas,
            "fuentes": {
                "estadoYPrecio": fuente_estado,
                "superficies": DOC_XLSX if x else DOC_PDF,
                "plano": "BROCHURE-CITY-HOMES.pdf — planta del piso correspondiente",
            },
            "historico": {
                "listaOficialPdf": {
                    "estado": estado_pdf[0], "precio": estado_pdf[1], "documento": DOC_PDF,
                } if p else None,
                "precioPlanillaAnterior": interpretar_precio((x or {}).get("precioPrevio"))[1] if x else None,
            },
        })

    garajes = []
    for numero in sorted(set(pdf_garajes) | set(xls_garajes), key=int):
        estado, monto = interpretar_precio(xls_garajes.get(numero, pdf_garajes.get(numero)))
        estado_pdf = interpretar_precio(pdf_garajes.get(numero))
        if estado_pdf[0] != estado and numero in pdf_garajes and numero in xls_garajes:
            discrepancias.append({
                "tipo": "estado", "unidad": f"Garaje {numero}",
                "listaOficialPdf": estado_pdf[0], "planillaVigente": estado,
            })
        garajes.append({
            "numero": numero, "estado": estado, "precio": monto, "moneda": "USD",
            "fuente": DOC_XLSX if numero in xls_garajes else DOC_PDF,
        })

    salida = {
        "generado": "scripts/construir_datos.py",
        "fuentes": {
            "listaOficialPdf": DOC_PDF,
            "planillaVigente": DOC_XLSX,
            "brochure": "BROCHURE-CITY-HOMES.pdf (Drive > City Homes)",
        },
        "reglaDePrecedencia": (
            "Ante diferencias entre documentos vale el más reciente (planilla del 23/10/2026). "
            "La lectura del PDF oficial se conserva en historico.listaOficialPdf y la diferencia "
            "queda listada en discrepancias."
        ),
        "unidades": unidades,
        "garajes": garajes,
        "discrepancias": discrepancias,
    }

    DESTINO.mkdir(parents=True, exist_ok=True)
    (DESTINO / "unidades.json").write_text(
        json.dumps(salida, ensure_ascii=False, indent=1), encoding="utf-8"
    )

    # Tabla maestra en CSV, para revisar o compartir fuera de la web.
    with (FUENTE / "tabla-maestra-unidades.csv").open("w", encoding="utf-8-sig", newline="") as f:
        escritor = csv.writer(f, delimiter=";")
        escritor.writerow([
            "unidad", "piso", "nivel", "tipologia", "dormitorios", "banos",
            "area_habitable", "terrazas_uso_exclusivo", "sub_total", "areas_comunes",
            "area_total", "precio_usd", "estado", "precio_lista_pdf", "estado_lista_pdf",
            "observaciones",
        ])
        for u in unidades:
            s, h = u["superficies"], u["historico"]["listaOficialPdf"] or {}
            escritor.writerow([
                u["unidad"], u["piso"], u["nivel"], u["tipologia"], u["dormitorios"], u["banos"],
                s["habitable"], s["terrazasUsoExclusivo"], s["subTotal"], s["areasComunes"],
                s["total"], u["precio"], u["estado"], h.get("precio"), h.get("estado"),
                " | ".join(u["observaciones"]),
            ])

    por_estado = {}
    for u in unidades:
        por_estado[u["estado"]] = por_estado.get(u["estado"], 0) + 1
    print(f"unidades: {len(unidades)}  {por_estado}")
    print(f"garajes:  {len(garajes)}")
    print(f"discrepancias entre documentos: {len(discrepancias)}")
    for d in discrepancias:
        print("   ", d)


if __name__ == "__main__":
    construir()
